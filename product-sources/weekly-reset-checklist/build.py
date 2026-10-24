"""Build the Weekly Reset Checklist workbook ($15 product).

Run:  python3 product-sources/weekly-reset-checklist/build.py
      (then recalculate the .xlsx in LibreOffice so formulas carry values, and
      run it again with --zip)
Out:  product-sources/dist/weekly-reset-checklist.zip  (upload to R2 as products/weekly-reset-checklist.zip)

The reset steps and the priority score (money x 2 + urgency - effort) match the
Digital System web app in src/data/appContent.ts and src/scripts/buzzyApp.ts.
"""

import os
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.join(HERE, "..", "dist")
XLSX = os.path.join(DIST, "Buzzyfly Weekly Reset Checklist.xlsx")
ZIP = os.path.join(DIST, "weekly-reset-checklist.zip")

FONT = "Arial"
BRAND = "1F4D3A"
INPUT_FILL = PatternFill("solid", fgColor="FFF6D5")
HEAD_FILL = PatternFill("solid", fgColor=BRAND)
TOP_FILL = PatternFill("solid", fgColor="CFE9DC")
LINE = Side(style="thin", color="E2DED3")
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

# Same steps and minutes as RESET_STEPS in src/data/appContent.ts.
RESET_STEPS = [
    (3, "Check the money", "Open invoices and anything due to be paid this week."),
    (5, "Send this week's follow-ups", "Work through everyone you owe a reply or a nudge."),
    (2, "Move each new client forward", "Tick off or send any onboarding step that's due."),
    (5, "List the week's tasks and score them", "Use the Priority Scorer tab. The top score is your #1."),
    (3, "Block time for your #1", "Put it in your calendar before anything else."),
    (2, "Drop or delay one thing", "Pick the task that can wait and tell anyone affected."),
]

TASK_ROWS = 15
LOG_WEEKS = 52


def f(bold=False, size=10, color="1D211F", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=16, color=BRAND)
    ws["A2"] = sub
    ws["A2"].font = f(italic=True, color="5F655F")


def header(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX


def inp(cell, value=None, fmt=None):
    """Mark a cell the buyer fills in: blue text on light yellow."""
    if value is not None:
        cell.value = value
    cell.font = f(color="0000FF")
    cell.fill = INPUT_FILL
    cell.border = BOX
    if fmt:
        cell.number_format = fmt


def calc(cell, formula, fmt=None, bold=False):
    cell.value = formula
    cell.font = f(bold=bold)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt


def widths(ws, cols):
    for letter, w in cols.items():
        ws.column_dimensions[letter].width = w


def build():
    wb = Workbook()
    yes_no = '"Yes,No"'

    # ---- Start Here -------------------------------------------------------
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "Buzzyfly Weekly Reset Checklist", "20 minutes on Monday. A quick check on Thursday. A close on Friday.")
    rows = [
        ("How it works", None),
        ("1. Monday", "Open the Monday Reset tab. Type the date, then work down the six steps and mark each one Yes."),
        ("2. Score your tasks", "On the Priority Scorer tab, list this week's tasks and score each one. The top score is your #1."),
        ("3. Thursday", "Open the Thursday Check-in tab. Five minutes: is your #1 moving, and what can you drop?"),
        ("4. Friday", "Open the Friday Close tab. Write down what got done and what came in, then copy the totals into the 52-Week Log."),
        ("5. Next Monday", "Clear the yellow cells on the Monday, Scorer, Thursday and Friday tabs, and start again. The log keeps your history."),
        (None, None),
        ("Which cells to edit", None),
        ("Yellow cells, blue text", "These are yours to fill in."),
        ("White cells", "These work themselves out. Don't type over them."),
        ("Example rows", "The first row on the Priority Scorer and the 52-Week Log is an example. Type over it or delete it."),
        (None, None),
        ("Help", "Email support@buzzyfly.com"),
    ]
    r = 4
    for label, text in rows:
        if label and text is None:
            ws.cell(row=r, column=1, value=label).font = f(bold=True, size=12, color=BRAND)
        elif label:
            ws.cell(row=r, column=1, value=label).font = f(bold=True)
            c = ws.cell(row=r, column=2, value=text)
            c.font = f()
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    inp(ws["A20"], "Example input")
    ws["B20"] = "← an input cell looks like this"
    ws["B20"].font = f(italic=True, color="5F655F")
    widths(ws, {"A": 24, "B": 90})

    # ---- Monday Reset -----------------------------------------------------
    ws = wb.create_sheet("Monday Reset")
    title(ws, "Monday Reset, 20 minutes", "Work down the list. Mark each step Yes when it's done.")
    ws["A4"] = "Week starting (Monday)"
    ws["A4"].font = f(bold=True)
    inp(ws["B4"], None, "ddd d mmm yyyy")
    ws["A5"] = "Week number"
    ws["A5"].font = f(bold=True)
    calc(ws["B5"], '=IF(B4="","",_xlfn.ISOWEEKNUM(B4))', "0")

    header(ws, 7, ["Step", "Minutes", "What to do", "Done?"])
    dv = DataValidation(type="list", formula1=yes_no, allow_blank=True)
    ws.add_data_validation(dv)
    for i, (minutes, label, detail) in enumerate(RESET_STEPS):
        row = 8 + i
        ws.cell(row=row, column=1, value=label).font = f(bold=True)
        ws.cell(row=row, column=2, value=minutes).font = f()
        c = ws.cell(row=row, column=3, value=detail)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).border = BOX
        inp(ws.cell(row=row, column=4))
        dv.add(ws.cell(row=row, column=4))
    last = 8 + len(RESET_STEPS) - 1
    ws.cell(row=last + 1, column=1, value="Total").font = f(bold=True)
    calc(ws.cell(row=last + 1, column=2), f"=SUM(B8:B{last})", "0", bold=True)
    ws.cell(row=last + 2, column=1, value="Steps done").font = f(bold=True)
    calc(ws.cell(row=last + 2, column=2), f'=COUNTIF(D8:D{last},"Yes")&" of {len(RESET_STEPS)}"', bold=True)

    m = last + 4
    ws.cell(row=m, column=1, value="Step 1: the money").font = f(bold=True, size=12, color=BRAND)
    money_rows = [
        ("Invoices still unpaid", "0"),
        ("Total still owed to you ($)", "$#,##0"),
        ("Bills you must pay this week ($)", "$#,##0"),
    ]
    for i, (label, fmt) in enumerate(money_rows, start=1):
        ws.cell(row=m + i, column=1, value=label).font = f()
        inp(ws.cell(row=m + i, column=2), None, fmt)

    ws.cell(row=m + 5, column=1, value="Your #1 this week").font = f(bold=True, size=12, color=BRAND)
    calc(
        ws.cell(row=m + 6, column=1),
        "=IFERROR(INDEX('Priority Scorer'!A6:A20,MATCH(1,'Priority Scorer'!G6:G20,0)),\"Score your tasks on the Priority Scorer tab\")",
        bold=True,
    )
    ws.merge_cells(start_row=m + 6, start_column=1, end_row=m + 6, end_column=3)
    ws.cell(row=m + 6, column=1).fill = TOP_FILL
    widths(ws, {"A": 34, "B": 16, "C": 58, "D": 10})

    # ---- Priority Scorer --------------------------------------------------
    ws = wb.create_sheet("Priority Scorer")
    title(ws, "Priority Scorer", "Score = money x 2 + urgency - effort. The highest score is your #1.")
    ws["A3"] = "Money 0-3: does it get you paid?   Urgency 0-3: is the deadline soon?   Effort 1-3: how big is it?"
    ws["A3"].font = f(color="5F655F")
    header(ws, 5, ["Task", "Money (0-3)", "Urgency (0-3)", "Effort (1-3)", "Score", "Notes", "Rank"])
    dv03 = DataValidation(type="whole", operator="between", formula1="0", formula2="3", allow_blank=True)
    dv13 = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=True)
    for d in (dv03, dv13):
        d.error = "Pick a whole number in the range shown in the column header."
        d.showErrorMessage = True
        ws.add_data_validation(d)
    first, end = 6, 6 + TASK_ROWS - 1
    example = ("Send the Acme proposal (example)", 3, 2, 1, "Due Wednesday")
    for row in range(first, end + 1):
        for col in (1, 2, 3, 4, 6):
            inp(ws.cell(row=row, column=col))
        dv03.add(ws.cell(row=row, column=2))
        dv03.add(ws.cell(row=row, column=3))
        dv13.add(ws.cell(row=row, column=4))
        calc(ws.cell(row=row, column=5), f'=IF(A{row}="","",B{row}*2+C{row}-D{row})', "0")
        # Ties go to the task listed first.
        calc(
            ws.cell(row=row, column=7),
            f'=IF(E{row}="","",COUNTIF($E${first}:$E${end},">"&E{row})+COUNTIF($E${first}:E{row},E{row}))',
            "0",
        )
    for col, value in enumerate(example, start=1):
        ws.cell(row=first, column=[1, 2, 3, 4, 6][col - 1], value=value)
    ws.conditional_formatting.add(
        f"A{first}:G{end}", FormulaRule(formula=[f"$G{first}=1"], fill=TOP_FILL, font=Font(name=FONT, bold=True))
    )
    ws.cell(row=end + 2, column=1, value="Your #1").font = f(bold=True, size=12, color=BRAND)
    calc(
        ws.cell(row=end + 3, column=1),
        f'=IFERROR(INDEX(A{first}:A{end},MATCH(1,G{first}:G{end},0)),"Add a task above")',
        bold=True,
    )
    ws.cell(row=end + 3, column=1).fill = TOP_FILL
    widths(ws, {"A": 44, "B": 12, "C": 13, "D": 12, "E": 9, "F": 34, "G": 8})
    ws.freeze_panes = "A6"

    # ---- Thursday Check-in ------------------------------------------------
    ws = wb.create_sheet("Thursday Check-in")
    title(ws, "Thursday Check-in, 5 minutes", "Catch the week before it slips.")
    ws["A4"] = "Your #1 this week"
    ws["A4"].font = f(bold=True)
    calc(ws["B4"], "='Monday Reset'!A" + str(last + 10))
    ws["B4"].fill = TOP_FILL
    questions = [
        ("Is your #1 done or on track?", True),
        ("If not, what's in the way?", False),
        ("One thing you can drop or move to next week", False),
        ("Who are you waiting on? Send them a nudge today.", False),
        ("Follow-ups sent so far this week", False),
    ]
    dv = DataValidation(type="list", formula1=yes_no, allow_blank=True)
    ws.add_data_validation(dv)
    for i, (q, is_yes_no) in enumerate(questions):
        row = 6 + i
        c = ws.cell(row=row, column=1, value=q)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        inp(ws.cell(row=row, column=2))
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if is_yes_no:
            dv.add(ws.cell(row=row, column=2))
    widths(ws, {"A": 46, "B": 60})

    # ---- Friday Close -----------------------------------------------------
    ws = wb.create_sheet("Friday Close")
    title(ws, "Friday Close, 10 minutes", "Write down the week, then copy the totals into the 52-Week Log.")
    ws["A4"] = "Week number"
    ws["A4"].font = f(bold=True)
    calc(ws["B4"], "='Monday Reset'!B5", "0")
    ws["A5"] = "Your #1 this week"
    ws["A5"].font = f(bold=True)
    calc(ws["B5"], "='Monday Reset'!A" + str(last + 10))
    close = [
        ("Did your #1 get done?", "list", None),
        ("Money that came in this week ($)", "num", "$#,##0"),
        ("Invoices sent this week", "num", "0"),
        ("Hours worked on client work", "num", "0.0"),
        ("Three things that got done", "text", None),
        ("What slipped, and why", "text", None),
        ("First task for next Monday", "text", None),
    ]
    dv = DataValidation(type="list", formula1=yes_no, allow_blank=True)
    ws.add_data_validation(dv)
    for i, (label, kind, fmt) in enumerate(close):
        row = 7 + i
        c = ws.cell(row=row, column=1, value=label)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        inp(ws.cell(row=row, column=2), None, fmt)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "list":
            dv.add(ws.cell(row=row, column=2))
    ws.cell(row=15, column=1, value="Now copy the week number, your #1, done, money and hours into the next empty row of the 52-Week Log.").font = f(
        italic=True, color="5F655F"
    )
    widths(ws, {"A": 40, "B": 60})

    # ---- 52-Week Log ------------------------------------------------------
    ws = wb.create_sheet("52-Week Log")
    title(ws, "52-Week Log", "One row per week. Your record of what mattered and what came in.")
    header(ws, 4, ["Week", "Your #1", "Done?", "Money in ($)", "Hours", "Notes", "Running total ($)"])
    first, end = 5, 5 + LOG_WEEKS - 1
    dv = DataValidation(type="list", formula1=yes_no, allow_blank=True)
    ws.add_data_validation(dv)
    for i in range(LOG_WEEKS):
        row = first + i
        calc(ws.cell(row=row, column=1), i + 1, "0")
        for col, fmt in ((2, None), (3, None), (4, "$#,##0"), (5, "0.0"), (6, None)):
            inp(ws.cell(row=row, column=col), None, fmt)
        dv.add(ws.cell(row=row, column=3))
        calc(ws.cell(row=row, column=7), f'=IF(D{row}="","",SUM(D${first}:D{row}))', "$#,##0")
    for col, value in ((2, "Send the Acme proposal (example)"), (3, "Yes"), (4, 1200), (5, 22), (6, "Example row. Type over it.")):
        ws.cell(row=first, column=col, value=value)

    s = end + 2
    summary = [
        ("Weeks logged", f'=COUNTA(B{first}:B{end})', "0"),
        ("Weeks your #1 got done", f'=COUNTIF(C{first}:C{end},"Yes")', "0"),
        ("Share of weeks your #1 got done", f'=IFERROR(B{s + 1}/B{s},0)', "0%"),
        ("Total money in ($)", f"=SUM(D{first}:D{end})", "$#,##0"),
        ("Average per logged week ($)", f'=IFERROR(B{s + 3}/COUNT(D{first}:D{end}),0)', "$#,##0"),
        ("Total hours", f"=SUM(E{first}:E{end})", "0.0"),
    ]
    for i, (label, formula, fmt) in enumerate(summary):
        ws.cell(row=s + i, column=1, value=label).font = f(bold=True)
        calc(ws.cell(row=s + i, column=2), formula, fmt, bold=True)
    widths(ws, {"A": 32, "B": 40, "C": 9, "D": 14, "E": 9, "F": 40, "G": 18})
    ws.freeze_panes = "A5"

    chart = BarChart()
    chart.title = "Money in by week"
    chart.y_axis.title = "$"
    chart.x_axis.title = "Week"
    chart.legend = None
    chart.add_data(Reference(ws, min_col=4, min_row=4, max_row=end), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=end))
    chart.height, chart.width = 8, 22
    ws.add_chart(chart, "I4")

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        # Print each tab one page wide.
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return wb


def main():
    os.makedirs(DIST, exist_ok=True)
    if "--zip" in sys.argv:
        with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(XLSX, os.path.basename(XLSX))
        print(ZIP)
        return
    build().save(XLSX)
    print(XLSX)


if __name__ == "__main__":
    main()
